"""
CRM de Búsqueda de Empleos — Sebastian Castro
==============================================
Ejecutar: streamlit run crm_empleos.py
Credenciales por defecto: admin / admin123
"""

import io
import json
import re
import subprocess
import sys
from dataclasses import asdict
from datetime import datetime, date, timedelta

import pandas as pd
import streamlit as st

# Instala los navegadores de Playwright si no están disponibles (Streamlit Cloud)
@st.cache_resource(show_spinner=False)
def _instalar_playwright():
    try:
        subprocess.run(
            [sys.executable, "-m", "playwright", "install", "chromium"],
            check=True, capture_output=True,
        )
    except Exception:
        pass

_instalar_playwright()
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import database as db

# ---------------------------------------------------------------------------
# Configuración de página
# ---------------------------------------------------------------------------

st.set_page_config(
    page_title="CRM Empleos",
    page_icon="💼",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------------------------------------------------------------------------
# CSS
# ---------------------------------------------------------------------------

st.markdown("""
<style>
.crm-header {
    background: linear-gradient(135deg,#1F4E79 0%,#2E75B6 100%);
    color:#fff; padding:1.1rem 1.4rem;
    border-radius:10px; margin-bottom:1rem;
}
.crm-header h1{margin:0;font-size:1.5rem}
.crm-header p{margin:0;opacity:.85;font-size:.88rem}

.metric-box{background:#f8f9fa;border:1px solid #dee2e6;
    border-radius:8px;padding:.75rem 1rem;text-align:center}
.metric-box .num{font-size:1.9rem;font-weight:700}
.metric-box .lbl{font-size:.78rem;color:#6c757d}

.badge{padding:3px 10px;border-radius:12px;font-size:.76rem;font-weight:600}
.badge-Pendiente{background:#6c757d;color:#fff}
.badge-Aplicada{background:#0d6efd;color:#fff}
.badge-AplicacionAvanzadaPendiente{background:#6610f2;color:#fff}
.badge-Entrevista{background:#fd7e14;color:#fff}
.badge-Oferta{background:#198754;color:#fff}
.badge-Descartada{background:#dc3545;color:#fff}
.badge-Alta{background:#198754;color:#fff}
.badge-Media{background:#ffc107;color:#000}
.badge-Baja{background:#dc3545;color:#fff}

.oferta-row{border-left:4px solid #2E75B6;padding:.5rem .8rem;
    margin:.4rem 0;background:#f8fbff;border-radius:0 6px 6px 0}

footer{visibility:hidden}
#MainMenu{visibility:hidden}
</style>
""", unsafe_allow_html=True)


# ---------------------------------------------------------------------------
# Session state helpers
# ---------------------------------------------------------------------------

def get_user() -> dict | None:
    return st.session_state.get("usuario")

def es_admin() -> bool:
    u = get_user()
    return u is not None and u.get("rol") == "admin"

def puede_editar() -> bool:
    u = get_user()
    return u is not None and u.get("rol") in ("admin", "usuario")

def logout():
    st.session_state.pop("usuario", None)
    st.rerun()


# ---------------------------------------------------------------------------
# Helpers UI
# ---------------------------------------------------------------------------

def badge(texto: str, clase: str) -> str:
    return f'<span class="badge badge-{clase}">{texto}</span>'

def metrica(col, numero, etiqueta, color="#1F4E79"):
    col.markdown(
        f'<div class="metric-box">'
        f'<div class="num" style="color:{color}">{numero}</div>'
        f'<div class="lbl">{etiqueta}</div></div>',
        unsafe_allow_html=True,
    )

_USD_COP = 4_200   # Tasa de conversión aproximada USD → COP

def _extraer_salario_cop(texto: str) -> tuple[int, int] | None:
    """
    Extrae (min, max) en COP de un texto de salario.
    Detecta moneda USD y convierte automáticamente.
    Retorna None si no hay valores numéricos reconocibles.
    """
    if not texto or texto.strip().lower() in ("no especificado", "", "—", "-"):
        return None
    tl = texto.lower()
    es_usd = any(k in tl for k in ("usd", "u.s.d", "dólar", "dolar", "us$"))

    nums_raw = re.findall(r"[\d]+(?:[.,][\d]+)*", texto)
    valores = []
    for n in nums_raw:
        try:
            # Determinar separador: si el último grupo tras punto/coma tiene 3 dígitos → miles
            n_clean = n
            if "." in n and "," in n:
                # Formato mixto: decidir cuál es decimal
                if n.index(",") < n.index("."):
                    n_clean = n.replace(",", "")          # 1,234.56
                else:
                    n_clean = n.replace(".", "").replace(",", ".")  # 1.234,56
                v = int(float(n_clean))
            elif "." in n:
                partes = n.split(".")
                if all(len(p) == 3 for p in partes[1:]):  # 1.234.567 → separador miles
                    v = int(n.replace(".", ""))
                else:
                    v = int(float(n))
            elif "," in n:
                partes = n.split(",")
                if all(len(p) == 3 for p in partes[1:]):  # 1,234,567 → separador miles
                    v = int(n.replace(",", ""))
                else:
                    v = int(float(n.replace(",", ".")))
            else:
                v = int(n)

            if es_usd:
                if v >= 100:                        # mínimo $100 USD
                    valores.append(v * _USD_COP)
            else:
                if v >= 100_000:                    # mínimo 100K COP
                    valores.append(v)
        except (ValueError, OverflowError):
            pass

    if not valores:
        return None
    return min(valores), max(valores)


def excel_en_memoria(df: pd.DataFrame, nombre_hoja: str = "Ofertas") -> bytes:
    """Genera un Excel en memoria y retorna bytes para st.download_button."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=nombre_hoja, index=False)
        ws = writer.sheets[nombre_hoja]
        # Encabezados en negrita y color
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E75B6")
            cell.alignment = Alignment(horizontal="center")
        # Ancho automático
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)
    buf.seek(0)
    return buf.read()


# ===========================================================================
# PANTALLA DE LOGIN
# ===========================================================================

def pantalla_login():
    col_c, col_form, col_d = st.columns([1.5, 2, 1.5])
    with col_form:
        st.markdown("<br><br>", unsafe_allow_html=True)
        st.markdown("""
        <div class="crm-header" style="text-align:center">
            <h1>💼 CRM Empleos</h1>
            <p>Sebastian Castro — Ingeniero de Datos Senior</p>
        </div>
        """, unsafe_allow_html=True)

        with st.form("login_form"):
            st.markdown("### Iniciar sesión")
            username = st.text_input("Usuario", placeholder="admin")
            password = st.text_input("Contraseña", type="password")
            submitted = st.form_submit_button("Ingresar", use_container_width=True, type="primary")

        if submitted:
            usuario = db.autenticar(username, password)
            if usuario:
                st.session_state["usuario"] = usuario
                st.success(f"Bienvenido, {usuario['nombre_completo']} 👋")
                st.rerun()
            else:
                st.error("Usuario o contraseña incorrectos.")

        st.caption("Credenciales por defecto: **admin** / **admin123**")


# ===========================================================================
# SIDEBAR
# ===========================================================================

def renderizar_sidebar() -> str:
    usuario = get_user()
    with st.sidebar:
        st.markdown(f"### 💼 CRM Empleos")
        st.markdown(
            f"👤 **{usuario['nombre_completo']}**  \n"
            f"Rol: `{usuario['rol']}`"
        )
        st.divider()

        opciones = ["📊 Dashboard", "📋 Mis Ofertas", "🔍 Buscar Ofertas",
                    "📈 Estadísticas", "📄 Reportes"]
        if es_admin():
            opciones += ["👥 Usuarios", "⚙️ Mi Perfil"]
        else:
            opciones += ["⚙️ Mi Perfil"]

        pagina = st.radio("Navegación", opciones, label_visibility="collapsed")

        st.divider()
        st.markdown("#### 🔎 Buscar por Job ID")
        job_id_input = st.number_input(
            "Job ID", min_value=0, step=1, value=0,
            help="Ingresa el ID numérico de la oferta para encontrarla directamente",
            label_visibility="collapsed",
        )
        if st.button("Ir a Job ID", use_container_width=True, disabled=(job_id_input == 0)):
            st.session_state["buscar_job_id"] = int(job_id_input)
            # Redirigir a Mis Ofertas
            st.session_state["pagina_forzada"] = "📋 Mis Ofertas"
            st.rerun()

        st.divider()
        st.markdown("#### Filtros")

        filtro_estado = st.multiselect(
            "Estado", db.ESTADOS,
            default=["Pendiente", "Aplicada", "Aplicacion Avanzada Pendiente", "Entrevista"],
        )
        filtro_fuente = st.multiselect(
            "Portal",
            ["LinkedIn", "Computrabajo", "Elempleo"],
            default=[], placeholder="Todos",
        )
        filtro_compat = st.multiselect(
            "Compatibilidad", ["Alta", "Media", "Baja"],
            default=[], placeholder="Todas",
        )
        filtro_modal = st.multiselect(
            "Modalidad", ["Remoto", "Híbrido"],
            default=[], placeholder="Todas",
        )
        col_f1, col_f2 = st.columns(2)
        fecha_desde = col_f1.date_input(
            "Desde", value=date.today()-timedelta(days=60),
            label_visibility="collapsed",
        )
        fecha_hasta = col_f2.date_input(
            "Hasta", value=date.today(),
            label_visibility="collapsed",
        )
        texto = st.text_input("🔎 Buscar texto", placeholder="empresa, tech…")

        st.divider()
        stats = db.estadisticas(usuario=get_user()["username"])
        st.caption(f"BD: **{stats['total']}** ofertas")
        if st.button("🚪 Cerrar sesión", use_container_width=True):
            logout()

        # Guardar filtros en session_state
        st.session_state["filtros"] = {
            "estado": filtro_estado or None,
            "fuentes": filtro_fuente or None,
            "nivel_compat": filtro_compat or None,
            "modalidad": filtro_modal or None,
            "fecha_desde": str(fecha_desde),
            "fecha_hasta": str(fecha_hasta),
            "texto": texto or None,
        }

    # Manejar redirección por Job ID
    if "pagina_forzada" in st.session_state:
        pagina = st.session_state.pop("pagina_forzada")
    return pagina


# ===========================================================================
# PÁGINA: Dashboard
# ===========================================================================

def pagina_dashboard():
    st.markdown("""
    <div class="crm-header">
        <h1>📊 Dashboard</h1>
        <p>Resumen de tu proceso de búsqueda de empleo</p>
    </div>
    """, unsafe_allow_html=True)

    stats = db.estadisticas(usuario=get_user()["username"])

    c1, c2, c3, c4, c5, c6 = st.columns(6)
    metrica(c1, stats["total"], "Total ofertas")
    metrica(c2, stats["por_estado"].get("Aplicada", 0), "Aplicadas", "#0d6efd")
    metrica(c3, stats["por_estado"].get("Entrevista", 0), "Entrevistas", "#fd7e14")
    metrica(c4, stats["por_estado"].get("Oferta recibida", 0), "Ofertas recibidas", "#198754")
    metrica(c5, stats["por_estado"].get("Descartada", 0), "Descartadas", "#dc3545")
    metrica(c6, f"{stats['prom_compat']}%", "Compat. promedio", "#6f42c1")

    st.markdown("---")
    col_izq, col_der = st.columns(2)

    with col_izq:
        st.markdown("#### Pipeline de aplicaciones")
        data_pipeline = {e: stats["por_estado"].get(e, 0) for e in db.ESTADOS}
        df_p = pd.DataFrame({"Estado": list(data_pipeline.keys()), "n": list(data_pipeline.values())})
        st.bar_chart(df_p.set_index("Estado"), color="#2E75B6")

    with col_der:
        st.markdown("#### Ofertas por portal")
        if stats["por_fuente"]:
            df_f = pd.DataFrame({"Portal": list(stats["por_fuente"].keys()), "n": list(stats["por_fuente"].values())})
            st.bar_chart(df_f.set_index("Portal"), color="#198754")
        else:
            st.info("Sin datos. Ejecuta una búsqueda primero.")

    st.markdown("---")
    n1, n2, n3 = st.columns(3)
    metrica(n1, stats["por_nivel"].get("Alta", 0), "Alta compatibilidad (≥65%)", "#198754")
    metrica(n2, stats["por_nivel"].get("Media", 0), "Media compatibilidad (35-64%)", "#ffc107")
    metrica(n3, stats["por_nivel"].get("Baja", 0), "Baja compatibilidad (<35%)", "#dc3545")

    st.markdown("---")
    st.markdown("#### 🏆 Top 10 mejores matches")
    top = db.obtener_ofertas(orden="compatibilidad_pct DESC", usuario=get_user()["username"])[:10]
    if top:
        # Cabecera
        h_id, h_t, h_c, h_e, h_l = st.columns([1, 4, 1.5, 1.8, 1])
        for col, txt in zip([h_id, h_t, h_c, h_e, h_l],
                            ["Job ID", "Oferta", "Compat.", "Estado", "Link"]):
            col.markdown(f"**{txt}**")

        for o in top:
            nivel_clase  = o["nivel_compat"]
            estado_clase = o["estado"].replace(" ", "")
            col_id, col_t, col_c, col_e, col_l = st.columns([1, 4, 1.5, 1.8, 1])
            col_id.caption(f"`#{o['id']}`")
            col_t.markdown(f"**{o['titulo']}** — {o['empresa'] or '—'}")
            col_c.markdown(
                badge(f"{o['compatibilidad_pct']}% {nivel_clase}", nivel_clase),
                unsafe_allow_html=True,
            )
            col_e.markdown(
                badge(f"{db.EMOJI_ESTADO.get(o['estado'],'')} {o['estado']}", estado_clase),
                unsafe_allow_html=True,
            )
            if o["url"]:
                col_l.markdown(f"[🔗 Ver]({o['url']})")
            else:
                col_l.caption("—")
    else:
        st.info("Sin ofertas aún. Ve a **Buscar Ofertas**.")


# ===========================================================================
# Componente reutilizable: detalle de una oferta
# ===========================================================================

def _cb_guardar_estado(oferta_id: int, key: str):
    """Callback on_change del selectbox — guarda el estado al instante."""
    nuevo = st.session_state.get(key)
    if nuevo:
        db.actualizar_estado(oferta_id, nuevo)
        st.toast(f"✓ Estado guardado: {nuevo}", icon="✅")


def _renderizar_oferta(o: dict, key_prefix: str = ""):
    """Muestra el detalle de una oferta y permite editar su estado."""
    nivel_clase  = o["nivel_compat"]
    estado_clase = o["estado"].replace(" ", "")
    emoji_e      = db.EMOJI_ESTADO.get(o["estado"], "")

    col_izq, col_der = st.columns([3, 2])

    with col_izq:
        st.markdown(
            f"**Job ID:** `#{o['id']}`  \n"
            f"**Portal:** {o['fuente']}  \n"
            f"**Empresa:** {o['empresa'] or '—'}  \n"
            f"**Ubicación:** {o['ubicacion'] or '—'}  \n"
            f"**Modalidad:** {o['modalidad']}  \n"
            f"**Salario:** {o['salario']}  \n"
            f"**Publicado:** {o['fecha_publicacion'] or '—'}  \n"
            f"**Extraído:** {o['fecha_extraccion']}"
        )
        if o.get("tecnologias_encontradas"):
            techs = [t for t in o["tecnologias_encontradas"].split(", ") if t]
            st.markdown("**Tecnologías detectadas:**")
            st.markdown(" ".join(f"`{t}`" for t in techs))
        if o.get("descripcion"):
            st.caption("**Descripción:**")
            st.caption(o["descripcion"][:500] + ("…" if len(o["descripcion"]) > 500 else ""))
        if o.get("url"):
            st.markdown(f"[🔗 **Ver oferta completa**]({o['url']})")
        else:
            st.caption("URL no disponible para esta oferta")

    with col_der:
        st.markdown(
            f"**Compatibilidad:** "
            + badge(f"{o['compatibilidad_pct']}% — {o['nivel_compat']}", nivel_clase),
            unsafe_allow_html=True,
        )
        st.markdown(
            f"**Estado actual:** "
            + badge(f"{emoji_e} {o['estado']}", estado_clase),
            unsafe_allow_html=True,
        )
        st.markdown("")

        if puede_editar():
            key_estado = f"{key_prefix}sel_estado_{o['id']}"
            st.selectbox(
                "Cambiar estado",
                db.ESTADOS,
                index=db.ESTADOS.index(o["estado"]) if o["estado"] in db.ESTADOS else 0,
                key=key_estado,
                on_change=_cb_guardar_estado,
                args=(o["id"], key_estado),
            )
            notas_val = st.text_area(
                "Notas",
                value=o.get("notas") or "",
                key=f"{key_prefix}notas_{o['id']}",
                height=90,
                placeholder="Contacto, salario ofrecido, próximos pasos…",
            )
            if st.button("💾 Guardar notas", key=f"{key_prefix}btn_{o['id']}", type="secondary"):
                estado_actual = st.session_state.get(key_estado, o["estado"])
                db.actualizar_estado(o["id"], estado_actual, notas_val)
                st.toast("✓ Notas guardadas", icon="📝")
        else:
            st.info("Solo lectura (rol Visualizador)")

        if o.get("fecha_aplicacion"):
            st.caption(f"📤 Aplicada: {o['fecha_aplicacion']}")
        if o.get("fecha_actualizacion"):
            st.caption(f"🕐 Actualizado: {o['fecha_actualizacion']}")


# ===========================================================================
# PÁGINA: Mis Ofertas
# ===========================================================================

def pagina_mis_ofertas():
    st.markdown("""
    <div class="crm-header">
        <h1>📋 Mis Ofertas</h1>
        <p>Gestiona el estado de cada postulación</p>
    </div>
    """, unsafe_allow_html=True)

    # ── Búsqueda directa por Job ID ──────────────────────────────────────────
    job_id_buscado = st.session_state.pop("buscar_job_id", None)
    if job_id_buscado:
        oferta_directa = db.buscar_por_job_id(job_id_buscado, usuario=get_user()["username"])
        if oferta_directa:
            st.success(f"✓ Oferta encontrada — Job ID **#{job_id_buscado}**")
            _renderizar_oferta(oferta_directa, key_prefix="jid_")
            st.divider()
            st.markdown("---")
        else:
            st.error(f"No se encontró ninguna oferta con Job ID **#{job_id_buscado}**")

    # ── Lista filtrada ────────────────────────────────────────────────────────
    f = st.session_state.get("filtros", {})
    ofertas = db.obtener_ofertas(
        estado=f.get("estado"),
        fuentes=f.get("fuentes"),
        nivel_compat=f.get("nivel_compat"),
        modalidad=f.get("modalidad"),
        busqueda_texto=f.get("texto"),
        fecha_desde=f.get("fecha_desde"),
        fecha_hasta=f.get("fecha_hasta"),
        orden="compatibilidad_pct DESC",
        usuario=get_user()["username"],
    )

    col_info, col_exp = st.columns([5, 1])
    col_info.markdown(f"**{len(ofertas)} ofertas** con los filtros actuales")

    if ofertas:
        df_exp = pd.DataFrame(ofertas)
        cols_exp = ["id", "titulo", "empresa", "fuente", "modalidad", "salario",
                    "ubicacion", "compatibilidad_pct", "nivel_compat",
                    "tecnologias_encontradas", "estado",
                    "fecha_publicacion", "fecha_extraccion", "notas", "url"]
        df_exp = df_exp[[c for c in cols_exp if c in df_exp.columns]]
        df_exp.rename(columns={"id": "Job ID"}, inplace=True)
        excel_bytes = excel_en_memoria(df_exp, "Ofertas")
        col_exp.download_button(
            "⬇️ Excel",
            data=excel_bytes,
            file_name=f"ofertas_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    if not ofertas:
        st.info("Sin ofertas con los filtros seleccionados.")
        return

    st.divider()

    _ESTADO_INDICADOR = {
        "Descartada":                    "🔴",
        "Aplicada":                      "🟢",
        "Aplicacion Avanzada Pendiente": "🟣",
        "Oferta recibida":               "💚",
        "Entrevista":                    "🟡",
        "Pendiente":                     "⚪",
    }
    _ESTADO_BG = {
        "Descartada":                    "#fff5f5",
        "Aplicada":                      "#f0fff4",
        "Aplicacion Avanzada Pendiente": "#f3e8ff",
        "Oferta recibida":               "#d4edda",
        "Entrevista":                    "#fff8e1",
        "Pendiente":                     "#f8f9fa",
    }
    _ESTADO_BORDER = {
        "Descartada":                    "#dc3545",
        "Aplicada":                      "#198754",
        "Aplicacion Avanzada Pendiente": "#6610f2",
        "Oferta recibida":               "#198754",
        "Entrevista":                    "#fd7e14",
        "Pendiente":                     "#6c757d",
    }

    for idx, o in enumerate(ofertas):
        ind   = _ESTADO_INDICADOR.get(o["estado"], "⚪")
        bg    = _ESTADO_BG.get(o["estado"], "#f8f9fa")
        borde = _ESTADO_BORDER.get(o["estado"], "#6c757d")

        # Franja de color antes del expander
        st.markdown(
            f'<div style="border-left:5px solid {borde};background:{bg};'
            f'border-radius:0 6px 0 0;padding:2px 8px;margin-bottom:-8px;'
            f'font-size:0.78rem;color:{borde};font-weight:600">'
            f'{ind} {o["estado"].upper()}</div>',
            unsafe_allow_html=True,
        )
        with st.expander(
            f"[#{o['id']}] **{o['titulo']}**  |  "
            f"{o['empresa'] or '—'}  |  "
            f"{o['compatibilidad_pct']}% {o['nivel_compat']}  |  "
            f"{o['fuente']}",
            expanded=False,
        ):
            _renderizar_oferta(o, key_prefix=f"lst{idx}_")


# ===========================================================================
# Componente: tabla de resultados de búsqueda con descarga
# ===========================================================================

def _mostrar_tabla_resultados(ofertas_dict: list[dict], busqueda_id: int):
    """
    Muestra tabla de resultados.
    Prioriza los registros de la BD (que ya tienen Job ID asignado).
    """
    # Obtener registros reales desde la BD para tener el Job ID (solo del usuario actual)
    ofs_bd = db.ofertas_por_busqueda(busqueda_id, usuario=get_user()["username"])

    if ofs_bd:
        df_src = pd.DataFrame(ofs_bd)
        cols_tabla = ["id", "titulo", "empresa", "fuente", "modalidad",
                      "compatibilidad_pct", "nivel_compat", "salario",
                      "tecnologias_encontradas", "url"]
    else:
        # Fallback: usar el dict en memoria (sin Job ID aún)
        df_src = pd.DataFrame(ofertas_dict)
        cols_tabla = ["titulo", "empresa", "fuente", "modalidad",
                      "compatibilidad_pct", "nivel_compat", "salario",
                      "tecnologias_encontradas", "url"]

    df_vista = df_src[[c for c in cols_tabla if c in df_src.columns]].copy()
    df_vista.rename(columns={
        "id": "Job ID",
        "titulo": "Titulo", "empresa": "Empresa", "fuente": "Portal",
        "modalidad": "Modalidad", "compatibilidad_pct": "Compat %",
        "nivel_compat": "Nivel", "salario": "Salario",
        "tecnologias_encontradas": "Tecnologías", "url": "URL",
    }, inplace=True)

    st.dataframe(
        df_vista, use_container_width=True, hide_index=True,
        column_config={
            "Job ID":   st.column_config.NumberColumn("Job ID", format="%d"),
            "URL":      st.column_config.LinkColumn("URL", display_text="🔗 Ver"),
            "Compat %": st.column_config.NumberColumn(format="%d%%"),
        },
    )

    # Excel con Job IDs
    cols_xls = ["id", "titulo", "empresa", "fuente", "modalidad", "salario",
                "ubicacion", "compatibilidad_pct", "nivel_compat",
                "tecnologias_encontradas", "estado", "fecha_publicacion", "url"]
    df_xls = pd.DataFrame(ofs_bd) if ofs_bd else df_src
    df_xls = df_xls[[c for c in cols_xls if c in df_xls.columns]].copy()
    df_xls.rename(columns={
        "id": "Job ID", "compatibilidad_pct": "Compat%",
        "nivel_compat": "Nivel", "tecnologias_encontradas": "Tecnologías",
    }, inplace=True)

    excel_bytes = excel_en_memoria(df_xls, f"Busqueda_{busqueda_id}")
    st.download_button(
        f"⬇️ Descargar Excel de esta búsqueda (#{busqueda_id})",
        data=excel_bytes,
        file_name=f"busqueda_{busqueda_id}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key=f"dl_busqueda_{busqueda_id}",
    )


# ===========================================================================
# PÁGINA: Buscar Ofertas
# ===========================================================================

def pagina_buscar():
    st.markdown("""
    <div class="crm-header">
        <h1>🔍 Buscar Nuevas Ofertas</h1>
        <p>Configura parámetros y lanza la búsqueda en todos los portales</p>
    </div>
    """, unsafe_allow_html=True)

    if not puede_editar():
        st.warning("Necesitas rol de Usuario o Admin para ejecutar búsquedas.")
        return

    with st.form("form_busqueda"):
        c1, c2 = st.columns(2)
        with c1:
            keywords_raw = st.text_area(
                "Términos de búsqueda (uno por línea)",
                value="Ingeniero de Datos\nData Engineer\nAnalista de Datos\nAnalista de Base de Datos\nAnalytics Engineer",
                height=140,
            )
            modalidades = st.multiselect(
                "Modalidad", ["Remoto", "Híbrido"],
                default=["Remoto", "Híbrido"],
            )
        with c2:
            ubicacion = st.text_input("Ubicación", value="Colombia")
            max_res = st.slider("Máx. resultados por portal", 20, 150, 60, 10)
            PORTALES_DISPONIBLES = [
                "LinkedIn", "Computrabajo", "Elempleo",
            ]
            portales = st.multiselect(
                "Portales",
                PORTALES_DISPONIBLES,
                default=PORTALES_DISPONIBLES,
                help="Elempleo usa Playwright (navegador real, más lento). LinkedIn y Computrabajo son más rápidos.",
            )

        c3, c4, c5 = st.columns(3)
        dias = c3.selectbox("Publicadas en los últimos", [7, 14, 30, 60, 90], index=2,
                            format_func=lambda x: f"{x} días")
        sin_ingles = c4.checkbox("Solo sin requisito de inglés", value=True)

        st.markdown("**Rango de salario mensual (COP millones)**")
        sal_cols = st.columns([4, 1])
        with sal_cols[0]:
            sal_rango = st.slider(
                "Salario COP",
                min_value=0, max_value=30,
                value=(0, 30),
                step=1,
                format="$%dM",
                label_visibility="collapsed",
            )
        with sal_cols[1]:
            incluir_sin_sal = st.checkbox("Incluir sin salario", value=True,
                                          help="Las ofertas que no publican salario siempre se incluyen si esta opción está marcada")

        lanzar = st.form_submit_button("🚀 Lanzar búsqueda", use_container_width=True, type="primary")

    if lanzar:
        keywords = [k.strip() for k in keywords_raw.strip().split("\n") if k.strip()]
        config = {
            "keywords": keywords, "modalidades": modalidades,
            "ubicacion": ubicacion, "max_por_portal": max_res,
            "portales": portales, "dias": dias, "sin_ingles": sin_ingles,
        }

        # st.status evita el error insertBefore de Streamlit al mezclar
        # progress + warnings en el mismo render tree
        todas        = []
        logs_errores = []
        usuario_actual = get_user()["username"]

        with st.status("⏳ Ejecutando búsqueda en portales…", expanded=True) as status:
            try:
                from buscador_empleos import (
                    ScraperLinkedIn, ScraperComputrabajo, ScraperElempleo,
                    calcular_compatibilidad, get_session,
                )
                from playwright.sync_api import sync_playwright
                import buscador_empleos as be_mod

                original_kw = be_mod.KEYWORDS_BUSQUEDA
                be_mod.KEYWORDS_BUSQUEDA = keywords
                session = get_session()

                # ── Portales via requests (sin navegador) ──────────────────
                if "LinkedIn" in portales:
                    st.write("🔎 Buscando en **LinkedIn**…")
                    try:
                        li = ScraperLinkedIn().buscar(session, max_total=max_res)
                        todas.extend(li)
                        st.write(f"✅ LinkedIn — {len(li)} ofertas")
                    except Exception as e:
                        logs_errores.append(f"LinkedIn: {e}")
                        st.write(f"⚠️ LinkedIn — error: {e}")

                if "Computrabajo" in portales:
                    st.write("🔎 Buscando en **Computrabajo**…")
                    try:
                        ct = ScraperComputrabajo().buscar(session, max_total=max_res)
                        todas.extend(ct)
                        st.write(f"✅ Computrabajo — {len(ct)} ofertas")
                    except Exception as e:
                        logs_errores.append(f"Computrabajo: {e}")
                        st.write(f"⚠️ Computrabajo — error: {e}")

                # ── Portales con Playwright (solo Elempleo) ────────────────
                if "Elempleo" in portales:
                    st.write("🌐 Iniciando navegador para **Elempleo**…")
                    try:
                        with sync_playwright() as pw:
                            browser = pw.chromium.launch(
                                headless=True, args=["--no-sandbox"]
                            )
                            ctx = browser.new_context(locale="es-CO")
                            ctx.add_init_script(
                                "Object.defineProperty(navigator,'webdriver',{get:()=>undefined})"
                            )
                            p = ctx.new_page()
                            el = ScraperElempleo().buscar(p, max_total=max_res)
                            p.close()
                            browser.close()
                        todas.extend(el)
                        st.write(f"✅ Elempleo — {len(el)} ofertas")
                    except Exception as e:
                        logs_errores.append(f"Elempleo: {e}")
                        st.write(f"⚠️ Elempleo — error: {e}")

                be_mod.KEYWORDS_BUSQUEDA = original_kw

                st.write("⚙️ Calculando compatibilidad, filtrando por salario y guardando en BD…")
                ofertas_dict = []
                descartadas_salario = 0
                sal_min_cop = sal_rango[0] * 1_000_000
                sal_max_cop = sal_rango[1] * 1_000_000
                aplicar_filtro_sal = (sal_rango[0] > 0 or sal_rango[1] < 30)

                for o in todas:
                    pct, nivel, techs = calcular_compatibilidad(o)
                    o.compatibilidad_pct      = pct
                    o.nivel_compat            = nivel
                    o.tecnologias_encontradas = techs
                    d = asdict(o)

                    # Filtro de salario
                    if aplicar_filtro_sal:
                        rango_sal = _extraer_salario_cop(d.get("salario", ""))
                        if rango_sal is None:
                            # Sin salario → incluir solo si el usuario lo permite
                            if not incluir_sin_sal:
                                descartadas_salario += 1
                                continue
                        else:
                            # Con salario → verificar que intersecte el rango
                            if rango_sal[0] > sal_max_cop or rango_sal[1] < sal_min_cop:
                                descartadas_salario += 1
                                continue

                    ofertas_dict.append(d)

                if descartadas_salario:
                    st.write(f"🔍 Filtro salario: {descartadas_salario} ofertas excluidas por no cumplir el rango")

                bid    = db.registrar_busqueda(config, 0, len(todas), usuario=usuario_actual)
                nuevas, dup = db.insertar_ofertas(
                    ofertas_dict, creado_por=usuario_actual, busqueda_id=bid
                )
                with db.conectar() as _c:
                    _c.execute("UPDATE busquedas SET total_nuevas=? WHERE id=?", (nuevas, bid))
                    _c.commit()

                status.update(
                    label=f"✅ Búsqueda #{bid} completada — {len(todas)} encontradas, {nuevas} nuevas",
                    state="complete", expanded=False,
                )

            except Exception as e:
                import traceback
                status.update(label=f"❌ Error en la búsqueda: {e}", state="error")
                st.code(traceback.format_exc())
                return

        # Fuera del st.status — sin riesgo de insertBefore
        st.success(
            f"✅ **{len(todas)}** ofertas encontradas — "
            f"**{nuevas}** nuevas guardadas — "
            f"**{dup}** ya estaban en la BD  |  🔖 Búsqueda **#{bid}**"
        )
        if nuevas > 0:
            st.balloons()

        if ofertas_dict:
            ofertas_dict.sort(key=lambda x: x.get("compatibilidad_pct", 0), reverse=True)
            st.markdown(f"### Resultados de la búsqueda #{bid} — {len(ofertas_dict)} ofertas")
            _mostrar_tabla_resultados(ofertas_dict, bid)

    # ── Zona de peligro (solo admin) ─────────────────────────────────────────
    if es_admin():
        st.markdown("---")
        with st.expander("⚠️ Zona de peligro — Administrador"):
            st.warning(
                "Esta acción elimina **permanentemente** todas las vacantes de la base de datos. "
                "El historial de búsquedas se mantiene, pero los datos de ofertas no se pueden recuperar."
            )
            confirmar_txt = st.text_input(
                "Escribe **ELIMINAR** para confirmar",
                placeholder="ELIMINAR",
                key="confirm_delete_inventario",
            )
            st.markdown("---")
            st.markdown("**Limpiar duplicados**")
            st.caption("Conserva la oferta más antigua (menor Job ID) cuando el mismo título+empresa+portal aparece más de una vez.")
            if st.button("🧹 Eliminar duplicados", use_container_width=True, type="secondary"):
                n = db.limpiar_duplicados()
                st.toast(f"✓ {n} ofertas duplicadas eliminadas.", icon="🧹")

            st.markdown("---")
            st.markdown("**Re-escanear ofertas existentes en la BD**")
            st.caption("Analiza título y descripción de todas las ofertas activas y descarta las que no cumplen criterios.")

            col_r1, col_r2, col_r3 = st.columns(3)
            with col_r1:
                if st.button("🌐 Descartar con inglés", use_container_width=True):
                    from buscador_empleos import PALABRAS_INGLES
                    n = db.descartar_ofertas_ingles(PALABRAS_INGLES)
                    st.toast(f"✓ {n} ofertas descartadas por inglés.", icon="🌐")
            with col_r2:
                if st.button("🏢 Descartar presenciales", use_container_width=True):
                    from buscador_empleos import PALABRAS_PRESENCIAL
                    n = db.descartar_presenciales(PALABRAS_PRESENCIAL)
                    st.toast(f"✓ {n} ofertas descartadas por ser presenciales.", icon="🏢")
            with col_r3:
                if st.button("🚫 Descartar BairesDev", use_container_width=True):
                    n = db.descartar_empresa("bairesdev")
                    st.toast(f"✓ {n} ofertas de BairesDev descartadas.", icon="🚫")

            st.markdown("---")
            if st.button("🗑️ Eliminar todo el inventario", type="primary",
                         disabled=(confirmar_txt.strip().upper() != "ELIMINAR")):
                n = db.eliminar_todas_ofertas()
                st.session_state["confirm_delete_inventario"] = ""
                st.toast(f"✓ Se eliminaron {n} vacantes de la base de datos.", icon="🗑️")

    # Historial de búsquedas con descarga por fila
    st.markdown("---")
    st.markdown("#### Historial de búsquedas")
    historial = db.historial_busquedas(15, usuario=get_user()["username"])
    if historial:
        # Cabecera
        h1, h2, h3, h4, h5, h6 = st.columns([2.2, 3.5, 1.2, 1.2, 1.5, 1.8])
        for col, txt in zip([h1,h2,h3,h4,h5,h6],
                            ["Fecha","Keywords","Encontradas","Nuevas","Usuario","Descargar"]):
            col.markdown(f"**{txt}**")
        st.divider()

        for h in historial:
            kw_txt = ", ".join(json.loads(h.get("parametros") or "{}").get("keywords", [])[:3])
            kw_txt = (kw_txt[:45] + "…") if len(kw_txt) > 45 else kw_txt
            c1, c2, c3, c4, c5, c6 = st.columns([2.2, 3.5, 1.2, 1.2, 1.5, 1.8])
            c1.caption(h.get("fecha","—"))
            c2.caption(kw_txt or "—")
            c3.caption(str(h.get("total_encontradas", 0)))
            c4.caption(str(h.get("total_nuevas", 0)))
            c5.caption(h.get("ejecutado_por","sistema") or "sistema")

            # Botón de descarga para cada búsqueda
            bid_h = h.get("id")
            if bid_h:
                ofs_h = db.ofertas_por_busqueda(bid_h, usuario=get_user()["username"])
                if ofs_h:
                    df_dl = pd.DataFrame(ofs_h)
                    cols_dl = ["id","titulo","empresa","fuente","modalidad","salario",
                               "compatibilidad_pct","nivel_compat","tecnologias_encontradas",
                               "estado","fecha_publicacion","url"]
                    df_dl = df_dl[[c for c in cols_dl if c in df_dl.columns]]
                    df_dl.rename(columns={"id":"Job ID","compatibilidad_pct":"Compat%",
                                          "nivel_compat":"Nivel","tecnologias_encontradas":"Tecnologías"},
                                 inplace=True)
                    xls = excel_en_memoria(df_dl, f"Busqueda_{bid_h}")
                    c6.download_button(
                        f"⬇️ Excel",
                        data=xls,
                        file_name=f"busqueda_{bid_h}_{h.get('fecha','')[:10]}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key=f"dl_hist_{bid_h}",
                    )
                else:
                    c6.caption("Sin datos")
            else:
                c6.caption("—")
    else:
        st.info("Sin historial de búsquedas.")


# ===========================================================================
# PÁGINA: Estadísticas
# ===========================================================================

def pagina_estadisticas():
    st.markdown("""
    <div class="crm-header">
        <h1>📈 Estadísticas</h1>
        <p>Analítica de tu proceso de búsqueda</p>
    </div>
    """, unsafe_allow_html=True)

    stats = db.estadisticas(usuario=get_user()["username"])
    todas = db.obtener_ofertas(usuario=get_user()["username"])

    if not todas:
        st.info("Sin datos. Ejecuta una búsqueda primero.")
        return

    df = pd.DataFrame(todas)

    c1, c2 = st.columns(2)
    with c1:
        st.markdown("#### Compatibilidad")
        vc = df["nivel_compat"].value_counts()
        st.bar_chart(vc, color="#2E75B6")
    with c2:
        st.markdown("#### Estado del pipeline")
        vc2 = df["estado"].value_counts()
        st.bar_chart(vc2, color="#198754")

    c3, c4 = st.columns(2)
    with c3:
        st.markdown("#### Por portal")
        st.bar_chart(df["fuente"].value_counts(), color="#fd7e14")
    with c4:
        st.markdown("#### Modalidad")
        st.bar_chart(df["modalidad"].value_counts(), color="#6f42c1")

    st.markdown("---")
    st.markdown("#### Tabla completa de ofertas")

    # Job ID siempre como primera columna
    cols = ["id", "titulo", "empresa", "fuente", "modalidad",
            "compatibilidad_pct", "nivel_compat", "estado", "salario",
            "tecnologias_encontradas", "fecha_publicacion",
            "fecha_extraccion", "url"]
    df_vista = df[[c for c in cols if c in df.columns]].copy()
    df_vista.rename(columns={
        "id": "Job ID",
        "titulo": "Titulo", "empresa": "Empresa", "fuente": "Portal",
        "modalidad": "Modalidad", "compatibilidad_pct": "Compat%",
        "nivel_compat": "Nivel", "estado": "Estado", "salario": "Salario",
        "tecnologias_encontradas": "Tecnologías",
        "fecha_publicacion": "Publicado", "fecha_extraccion": "Extraído",
        "url": "URL",
    }, inplace=True)

    st.dataframe(
        df_vista, use_container_width=True, hide_index=True,
        column_config={
            "Job ID":  st.column_config.NumberColumn("Job ID", format="%d"),
            "URL":     st.column_config.LinkColumn("URL", display_text="🔗 Ver"),
            "Compat%": st.column_config.NumberColumn(format="%d%%"),
        },
    )

    excel_bytes = excel_en_memoria(df_vista, "Estadísticas")
    st.download_button(
        "⬇️ Exportar todo a Excel",
        data=excel_bytes,
        file_name=f"estadisticas_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ===========================================================================
# PÁGINA: Reportes
# ===========================================================================

def pagina_reportes():
    st.markdown("""
    <div class="crm-header">
        <h1>📄 Reportes</h1>
        <p>Genera reportes personalizados con los filtros que necesites y descárgalos en Excel</p>
    </div>
    """, unsafe_allow_html=True)

    todas = db.obtener_ofertas(usuario=get_user()["username"])
    if not todas:
        st.info("Sin datos. Ejecuta una búsqueda primero.")
        return

    st.markdown("#### Configura tu reporte")

    # ── Fila 1: Estados y Compatibilidad ─────────────────────────────────────
    col_a, col_b = st.columns(2)
    with col_a:
        st.markdown("**Estados a INCLUIR**")
        estados_inc = st.multiselect(
            "Estados incluir",
            db.ESTADOS,
            default=[e for e in db.ESTADOS if e != "Descartada"],
            label_visibility="collapsed",
            key="rep_estados",
        )
    with col_b:
        st.markdown("**Compatibilidad**")
        compat_inc = st.multiselect(
            "Compatibilidad",
            ["Alta", "Media", "Baja"],
            default=["Alta", "Media", "Baja"],
            label_visibility="collapsed",
            key="rep_compat",
        )

    # ── Fila 2: Portal y Modalidad ────────────────────────────────────────────
    col_c, col_d = st.columns(2)
    with col_c:
        st.markdown("**Portal**")
        fuentes_inc = st.multiselect(
            "Portal",
            ["LinkedIn", "Computrabajo", "Elempleo"],
            default=[],
            placeholder="Todos",
            label_visibility="collapsed",
            key="rep_fuente",
        )
    with col_d:
        st.markdown("**Modalidad**")
        modal_inc = st.multiselect(
            "Modalidad",
            ["Remoto", "Híbrido"],
            default=[],
            placeholder="Todas",
            label_visibility="collapsed",
            key="rep_modal",
        )

    # ── Fila 3: Fechas y texto libre ──────────────────────────────────────────
    col_e, col_f, col_g = st.columns([2, 2, 3])
    with col_e:
        st.markdown("**Fecha extracción desde**")
        rep_desde = st.date_input(
            "Desde",
            value=date.today() - timedelta(days=90),
            label_visibility="collapsed",
            key="rep_desde",
        )
    with col_f:
        st.markdown("**Hasta**")
        rep_hasta = st.date_input(
            "Hasta",
            value=date.today(),
            label_visibility="collapsed",
            key="rep_hasta",
        )
    with col_g:
        st.markdown("**Buscar texto**")
        rep_texto = st.text_input(
            "Texto",
            placeholder="empresa, tecnología, título…",
            label_visibility="collapsed",
            key="rep_texto",
        )

    # ── Fila 4: Rango de salario ──────────────────────────────────────────────
    col_sal, col_sin_sal = st.columns([4, 2])
    with col_sal:
        st.markdown("**Rango de salario mensual (COP millones)**")
        rep_sal = st.slider(
            "Salario reporte",
            min_value=0, max_value=30,
            value=(0, 30),
            step=1,
            format="$%dM",
            label_visibility="collapsed",
            key="rep_sal",
        )
    with col_sin_sal:
        st.markdown("**Opciones**")
        rep_sin_sal = st.checkbox(
            "Incluir ofertas sin salario especificado",
            value=True,
            key="rep_sin_sal",
        )

    # ── Columnas a exportar ───────────────────────────────────────────────────
    st.markdown("**Columnas a incluir en el reporte**")
    todas_cols = {
        "Job ID": "id", "Título": "titulo", "Empresa": "empresa",
        "Portal": "fuente", "Modalidad": "modalidad", "Salario": "salario",
        "Ubicación": "ubicacion", "Compat%": "compatibilidad_pct",
        "Nivel": "nivel_compat", "Estado": "estado",
        "Tecnologías": "tecnologias_encontradas", "Publicado": "fecha_publicacion",
        "Extraído": "fecha_extraccion", "Aplicado": "fecha_aplicacion",
        "Actualizado": "fecha_actualizacion", "Notas": "notas", "URL": "url",
    }
    cols_sel = st.multiselect(
        "Columnas",
        list(todas_cols.keys()),
        default=["Job ID", "Título", "Empresa", "Portal", "Modalidad",
                 "Salario", "Compat%", "Nivel", "Estado", "Tecnologías",
                 "Publicado", "URL"],
        label_visibility="collapsed",
        key="rep_cols",
    )

    # ── Aplicar filtros ───────────────────────────────────────────────────────
    st.markdown("---")
    df_all = pd.DataFrame(todas)

    # Estado
    if estados_inc:
        df_all = df_all[df_all["estado"].isin(estados_inc)]
    # Compatibilidad
    if compat_inc and len(compat_inc) < 3:
        df_all = df_all[df_all["nivel_compat"].isin(compat_inc)]
    # Portal
    if fuentes_inc:
        df_all = df_all[df_all["fuente"].isin(fuentes_inc)]
    # Modalidad
    if modal_inc:
        df_all = df_all[df_all["modalidad"].isin(modal_inc)]
    # Fechas — fillna protege contra filas con fecha_extraccion NULL
    fechas = df_all["fecha_extraccion"].fillna("1970-01-01")
    df_all = df_all[(fechas >= str(rep_desde)) & (fechas <= str(rep_hasta) + " 23:59")]
    # Texto libre
    if rep_texto:
        t = rep_texto.lower()
        mask = (
            df_all["titulo"].fillna("").str.lower().str.contains(t, na=False) |
            df_all["empresa"].fillna("").str.lower().str.contains(t, na=False) |
            df_all["tecnologias_encontradas"].fillna("").str.lower().str.contains(t, na=False)
        )
        df_all = df_all[mask]

    # Salario
    aplicar_sal = (rep_sal[0] > 0 or rep_sal[1] < 30)
    if aplicar_sal:
        sal_min_cop = rep_sal[0] * 1_000_000
        sal_max_cop = rep_sal[1] * 1_000_000
        def _pasa_salario(row):
            rango = _extraer_salario_cop(row.get("salario", ""))
            if rango is None:
                return rep_sin_sal
            return rango[0] <= sal_max_cop and rango[1] >= sal_min_cop
        df_all = df_all[df_all.apply(_pasa_salario, axis=1)]

    # ── Vista previa ──────────────────────────────────────────────────────────
    n_res = len(df_all)
    st.markdown(f"#### Vista previa — **{n_res}** ofertas")

    if n_res == 0:
        st.warning("Ninguna oferta cumple los filtros seleccionados. Ajusta los criterios.")
        return

    # Construir df de vista con las columnas seleccionadas
    cols_bd = [todas_cols[c] for c in cols_sel if todas_cols[c] in df_all.columns]
    df_vista = df_all[cols_bd].copy()
    rename_map = {v: k for k, v in todas_cols.items() if v in cols_bd}
    df_vista.rename(columns=rename_map, inplace=True)

    cfg = {}
    if "Job ID" in df_vista.columns:
        cfg["Job ID"] = st.column_config.NumberColumn("Job ID", format="%d")
    if "URL" in df_vista.columns:
        cfg["URL"] = st.column_config.LinkColumn("URL", display_text="🔗 Ver")
    if "Compat%" in df_vista.columns:
        cfg["Compat%"] = st.column_config.NumberColumn(format="%d%%")

    st.dataframe(df_vista, use_container_width=True, hide_index=True, column_config=cfg)

    # ── Resumen rápido ────────────────────────────────────────────────────────
    with st.expander("📊 Resumen del reporte"):
        r1, r2, r3, r4 = st.columns(4)
        metrica(r1, n_res, "Ofertas en el reporte")
        metrica(r2, df_all["estado"].value_counts().get("Aplicada", 0), "Aplicadas", "#0d6efd")
        metrica(r3, df_all["nivel_compat"].value_counts().get("Alta", 0), "Alta compat.", "#198754")
        metrica(r4, df_all["fuente"].nunique(), "Portales")

        col_g1, col_g2 = st.columns(2)
        with col_g1:
            try:
                vc_estado = df_all["estado"].fillna("Sin estado").value_counts()
                if not vc_estado.empty:
                    st.markdown("**Por estado**")
                    st.bar_chart(vc_estado, color="#2E75B6")
            except Exception:
                pass
        with col_g2:
            try:
                vc_fuente = df_all["fuente"].fillna("Sin portal").value_counts()
                if not vc_fuente.empty:
                    st.markdown("**Por portal**")
                    st.bar_chart(vc_fuente, color="#198754")
            except Exception:
                pass

    # ── Descarga Excel ────────────────────────────────────────────────────────
    nombre_archivo = (
        f"reporte_empleos_{datetime.now().strftime('%Y%m%d_%H%M')}"
        f"_{'_'.join(e[:3] for e in estados_inc) if estados_inc else 'todos'}.xlsx"
    )
    excel_bytes = excel_en_memoria(df_vista, "Reporte")
    st.download_button(
        f"⬇️ Descargar reporte en Excel ({n_res} ofertas)",
        data=excel_bytes,
        file_name=nombre_archivo,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
        key="dl_reporte",
    )


# ===========================================================================
# PÁGINA: Gestión de Usuarios (solo admin)
# ===========================================================================

def pagina_usuarios():
    if not es_admin():
        st.error("Acceso denegado. Solo administradores.")
        return

    st.markdown("""
    <div class="crm-header">
        <h1>👥 Gestión de Usuarios</h1>
        <p>Crear, editar y administrar usuarios del CRM</p>
    </div>
    """, unsafe_allow_html=True)

    tab_lista, tab_nuevo = st.tabs(["📋 Usuarios existentes", "➕ Crear usuario"])

    # ---- Tab: lista de usuarios ----
    with tab_lista:
        usuarios = db.obtener_usuarios()
        if not usuarios:
            st.info("Sin usuarios.")
        else:
            for u in usuarios:
                with st.expander(
                    f"{'🟢' if u['activo'] else '🔴'} **{u['username']}** — "
                    f"{u['nombre_completo']} | Rol: `{u['rol']}`"
                ):
                    col_datos, col_acciones = st.columns([3, 2])

                    with col_datos:
                        st.markdown(
                            f"**Email:** {u['email'] or '—'}  \n"
                            f"**Creado:** {u['fecha_creacion']}  \n"
                            f"**Último acceso:** {u['ultimo_acceso'] or '—'}  \n"
                            f"**Activo:** {'Sí' if u['activo'] else 'No'}"
                        )

                    with col_acciones:
                        with st.form(f"form_editar_{u['id']}"):
                            nuevo_nombre = st.text_input("Nombre", value=u["nombre_completo"] or "")
                            nuevo_email  = st.text_input("Email", value=u["email"] or "")
                            nuevo_rol    = st.selectbox(
                                "Rol", db.ROLES,
                                index=db.ROLES.index(u["rol"]) if u["rol"] in db.ROLES else 1,
                            )
                            nuevo_activo = st.checkbox("Activo", value=bool(u["activo"]))
                            nueva_pwd    = st.text_input("Nueva contraseña (dejar vacío = sin cambio)", type="password")

                            col_g, col_e = st.columns(2)
                            guardar  = col_g.form_submit_button("💾 Guardar", type="primary")
                            eliminar = col_e.form_submit_button(
                                "🗑️ Eliminar",
                                disabled=(u["username"] == "admin"),
                            )

                        if guardar:
                            db.actualizar_usuario(u["id"], nuevo_nombre, nuevo_email, nuevo_rol, nuevo_activo)
                            if nueva_pwd:
                                db.cambiar_password(u["id"], nueva_pwd)
                            st.toast("✓ Usuario actualizado", icon="✅")

                        if eliminar and u["username"] != "admin":
                            db.eliminar_usuario(u["id"])
                            st.toast("✓ Usuario eliminado", icon="🗑️")

    # ---- Tab: crear usuario ----
    with tab_nuevo:
        with st.form("form_crear_usuario"):
            st.markdown("### Nuevo usuario")
            c1, c2 = st.columns(2)
            nuevo_user  = c1.text_input("Username *")
            nueva_pwd2  = c2.text_input("Contraseña *", type="password")
            nuevo_nom   = c1.text_input("Nombre completo *")
            nuevo_email = c2.text_input("Email")
            nuevo_rol2  = st.selectbox("Rol", db.ROLES, index=1)
            crear = st.form_submit_button("✅ Crear usuario", type="primary")

        if crear:
            if not nuevo_user or not nueva_pwd2 or not nuevo_nom:
                st.error("Username, contraseña y nombre son obligatorios.")
            else:
                ok = db.crear_usuario(nuevo_user, nueva_pwd2, nuevo_nom, nuevo_email, nuevo_rol2)
                if ok:
                    st.toast(f"✓ Usuario {nuevo_user} creado con rol {nuevo_rol2}", icon="✅")
                else:
                    st.error("El username ya existe.")


# ===========================================================================
# PÁGINA: Mi Perfil
# ===========================================================================

def pagina_perfil():
    st.markdown("""
    <div class="crm-header">
        <h1>⚙️ Mi Perfil</h1>
        <p>Gestiona tu información y contraseña</p>
    </div>
    """, unsafe_allow_html=True)

    u = get_user()
    col_info, col_pwd = st.columns(2)

    with col_info:
        st.markdown("#### Información de cuenta")
        st.markdown(
            f"**Usuario:** `{u['username']}`  \n"
            f"**Nombre:** {u['nombre_completo']}  \n"
            f"**Email:** {u['email'] or '—'}  \n"
            f"**Rol:** `{u['rol']}`  \n"
            f"**Último acceso:** {u['ultimo_acceso'] or '—'}"
        )

    with col_pwd:
        st.markdown("#### Cambiar contraseña")
        with st.form("form_pwd"):
            pwd_actual  = st.text_input("Contraseña actual", type="password")
            pwd_nueva   = st.text_input("Nueva contraseña", type="password")
            pwd_confirm = st.text_input("Confirmar nueva contraseña", type="password")
            cambiar = st.form_submit_button("🔒 Cambiar contraseña", type="primary")

        if cambiar:
            verificado = db.autenticar(u["username"], pwd_actual)
            if not verificado:
                st.error("Contraseña actual incorrecta.")
            elif pwd_nueva != pwd_confirm:
                st.error("Las contraseñas nuevas no coinciden.")
            elif len(pwd_nueva) < 6:
                st.error("La contraseña debe tener al menos 6 caracteres.")
            else:
                db.cambiar_password(u["id"], pwd_nueva)
                st.success("Contraseña actualizada ✓")


# ===========================================================================
# ENTRADA PRINCIPAL
# ===========================================================================

def main():
    # Si no hay sesión activa → login
    if not get_user():
        pantalla_login()
        return

    pagina = renderizar_sidebar()

    if pagina == "📊 Dashboard":
        pagina_dashboard()
    elif pagina == "📋 Mis Ofertas":
        pagina_mis_ofertas()
    elif pagina == "🔍 Buscar Ofertas":
        pagina_buscar()
    elif pagina == "📈 Estadísticas":
        pagina_estadisticas()
    elif pagina == "📄 Reportes":
        pagina_reportes()
    elif pagina == "👥 Usuarios":
        pagina_usuarios()
    elif pagina == "⚙️ Mi Perfil":
        pagina_perfil()


if __name__ == "__main__":
    main()
